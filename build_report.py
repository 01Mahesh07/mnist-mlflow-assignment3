"""Fill the Word template with actual local experiment results and DagsHub screenshots.

Run: python build_report.py
This script intentionally leaves missing evidence marked; it does not invent results.
"""
from __future__ import annotations
import csv
import os
from pathlib import Path
from dotenv import load_dotenv
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "Assignment3_Report_Template.docx"
FINAL = ROOT / "Assignment3_Final_Report.docx"
RESULTS = ROOT / "outputs" / "experiment_results.csv"
SCREENSHOTS = {
    "{{SCREENSHOT_RUNS}}": "01_run_list.png",
    "{{SCREENSHOT_PARAMS}}": "02_logged_parameters.png",
    "{{SCREENSHOT_METRICS}}": "03_logged_metrics.png",
    "{{SCREENSHOT_COMPARISON}}": "04_run_comparison.png",
}
KEYS = ("baseline", "wider", "deeper", "sgd")


def replace_tokens_in_paragraph(p, replacements):
    # The template deliberately keeps placeholder tokens in single paragraphs.
    source = p.text
    target = source
    for token, replacement in replacements.items():
        target = target.replace(token, replacement)
    if target != source:
        p.text = target


def main():
    load_dotenv(ROOT / ".env")
    doc = Document(TEMPLATE)
    results = {}
    if RESULTS.exists():
        with RESULTS.open(newline="", encoding="utf-8") as f:
            results = {r["run_key"]: r for r in csv.DictReader(f) if r["run_key"] in KEYS}

    owner = os.getenv("DAGSHUB_OWNER", "").strip()
    repo = os.getenv("DAGSHUB_REPO", "").strip()
    github = os.getenv("GITHUB_URL", "").strip()
    has_dag = bool(owner and repo and not owner.startswith("YOUR_") and not repo.startswith("YOUR_"))
    has_git = bool(github and "YOUR_" not in github and github.startswith("https://github.com/"))
    replacements = {
        "{{GITHUB_URL}}": github if has_git else "PENDING: supply your actual GitHub project URL after pushing code",
        "{{DAGSHUB_URL}}": (f"https://dagshub.com/{owner}/{repo}" if has_dag else "PENDING: create/connect your DagsHub repository"),
        "{{MLFLOW_TRACKING_URI}}": (f"https://dagshub.com/{owner}/{repo}.mlflow" if has_dag else "PENDING: configure your DagsHub MLflow Tracking URI"),
        "{{RUN_STATUS}}": (
            "All four required configurations have local results; verify that the matching run IDs appear in DagsHub."
            if len(results) == 4 else
            f"INCOMPLETE: {len(results)} of four required configurations have local results. Run missing configurations."
        ),
    }
    for key in KEYS:
        row = results.get(key)
        for field in ("train_accuracy", "validation_accuracy", "train_loss", "validation_loss"):
            token = "{{" + f"{key}_{field}" + "}}"
            if row:
                value = float(row[field])
                replacements[token] = f"{value*100:.2f}%" if "accuracy" in field else f"{value:.4f}"
            else:
                replacements[token] = "PENDING"
        replacements["{{" + f"{key}_id" + "}}"] = row["run_id"] if row else "PENDING"

    if len(results) == 4:
        best = sorted(results.values(), key=lambda r:(-float(r["validation_accuracy"]),float(r["validation_loss"])))[0]
        replacements["{{BEST_MODEL}}"] = (
            f"{best['run_name']} was selected because its final-epoch validation accuracy "
            f"({float(best['validation_accuracy'])*100:.2f}%) was highest among the four tested models. "
            f"Its final validation loss was {float(best['validation_loss']):.4f}. "
            "If validation accuracy had tied, the lower validation loss would have been preferred. "
            "The independent MNIST test set was not used to select the model; validation results "
            "estimate performance on held-out data but do not guarantee test-set performance."
        )
    else:
        replacements["{{BEST_MODEL}}"] = (
            "PENDING: After all four runs finish, choose the largest final validation accuracy "
            "(use smaller final validation loss as a tie-breaker). Do not claim a winner before training."
        )

    for p in doc.paragraphs:
        if p.text.strip() in SCREENSHOTS:
            path = ROOT / "screenshots" / SCREENSHOTS[p.text.strip()]
            p.text = ""
            if path.is_file():
                p.add_run().add_picture(str(path), width=Inches(6.15))
            else:
                run = p.add_run(f"EVIDENCE NOT ADDED: Save your actual DagsHub capture to screenshots/{path.name} and regenerate this report.")
                run.bold = True
                run.font.color.rgb = RGBColor(150, 35, 35)
        else:
            replace_tokens_in_paragraph(p, replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_tokens_in_paragraph(p, replacements)
    doc.save(FINAL)
    missing_screens = [n for n in SCREENSHOTS.values() if not (ROOT / "screenshots" / n).is_file()]
    print(f"Created: {FINAL}")
    print(f"Recorded experiments: {len(results)}/4")
    if missing_screens:
        print("Still missing authentic screenshots:", ", ".join(missing_screens))
    if not has_git or not has_dag:
        print("Still missing your verified GitHub/DagsHub account details in .env")
    if len(results) == 4 and not missing_screens and has_git and has_dag:
        print("Evidence inserted. Verify URLs, run IDs, visual layout and accuracy before submission.")
    else:
        print("This report is a WORKING DRAFT, not an evidence-complete final submission.")


if __name__ == "__main__":
    main()
